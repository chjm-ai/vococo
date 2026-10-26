"""最终 PPTX 结构检查、私有交付与内置适配入口；不依赖旧生成器。"""
import importlib.util
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vococo/plugin/skills/editable-pptx"
spec = importlib.util.spec_from_file_location("pptx_check", SKILL / "scripts/check.py")
checker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(checker)


@pytest.fixture
def native_deck(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "原生文字"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "原生表格"
    output = tmp_path / "native.pptx"
    presentation.save(output)
    return output


def test_final_pptx_structure_check_does_not_claim_visual_success(native_deck):
    report = checker.inspect(native_deck)
    assert report["slide_count"] == 1 and not report["errors"]
    assert report["slides"][0]["texts"] == 2
    assert report["slides"][0]["tables"] == 1
    assert report["visual_check"] == "not_run"


def test_real_pptx_passes_private_artifact_safety_scan(native_deck, monkeypatch, isolated):
    from vococo.tools import artifacts
    monkeypatch.setattr(artifacts, "_roots", lambda *_: (native_deck.parent.resolve(),))
    name, data = artifacts._snapshot(str(native_deck), "web:pptx", str(native_deck.parent))
    assert name == native_deck.name and data == native_deck.read_bytes()


def test_internal_skill_available_even_with_skill_whitelist():
    from vococo.core.agent import _PLUGIN_SKILLS
    assert "vococo-internal:editable-pptx" in _PLUGIN_SKILLS
    text = (SKILL / "SKILL.md").read_text()
    assert "send_file" in text and "未验证实际 Office 排版" in text
