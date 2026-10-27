"""临时 Web 实例，从真实工具到聊天卡片，再实际点击浏览器下载。"""
import asyncio
import hashlib
from pathlib import Path

import pytest

from vococo import config
from vococo.core.agent import AgentReply
from vococo.gateway import clarify
from vococo.gateway.adapters.web import WebAdapter
from vococo.memory import session_store
from vococo.tools import builtin, danger
from test_library_integration import _browser_environment
from pptx import Presentation


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.mark.anyio
async def test_file_tool_card_history_and_browser_download(isolated, monkeypatch):
    monkeypatch.setattr(config, "WEB_HOST", "127.0.0.1")
    monkeypatch.setattr(config, "WEB_PORT", 0)
    monkeypatch.setattr(config, "WEB_AUTH_TOKEN", "file-browser-test-auth")
    monkeypatch.setattr(config, "VOICE_ENABLED", False)
    adapter = WebAdapter()
    root = isolated / "project"
    root.mkdir()
    file = root / "可编辑演示.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(file)
    content = file.read_bytes()

    async def drive_tool():
        message = await adapter._inbox.get()
        assert message.text == "deliver"
        key = config.resolve_session_key("web", "file-e2e")
        tid = session_store.start_turn(key, message.text)
        route = clarify.set_current(key, adapter, "file-e2e", tid)
        cwd = danger.set_cwd(str(root))
        try:
            result = await builtin.send_file.handler({"path": str(file), "title": "可编辑成果"})
            assert "已交付" in result["content"][0]["text"]
            file.unlink()
            finish = await adapter._inbox.get()
            assert finish.text == "finish"
            session_store.finish_turn(tid, "已生成可编辑成果。")
            await adapter.make_sink("file-e2e").done(AgentReply(text="已生成可编辑成果。", tool_calls=[],
                                                               cost_usd=0, is_error=False), tid)
        finally:
            danger.reset_cwd(cwd)
            clarify.reset_current(route)

    task = process = None
    try:
        await adapter._start_server()
        origin = f"http://127.0.0.1:{adapter._runner.addresses[0][1]}"
        task = asyncio.create_task(drive_tool())
        process = await asyncio.create_subprocess_exec(
            "node", str(Path(__file__).with_name("send_file_browser.cjs")), origin,
            config.WEB_AUTH_TOKEN, str(isolated), hashlib.sha256(content).hexdigest(),
            env=_browser_environment(), stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await asyncio.wait_for(process.communicate(), 120)
        if task.done():
            await task
        assert process.returncode == 0, stderr.decode() + stdout.decode()
        await asyncio.wait_for(task, 5)
        print(stdout.decode())
    finally:
        if process and process.returncode is None:
            process.kill()
            await process.wait()
        if task and not task.done():
            task.cancel()
            await asyncio.gather(task, return_exceptions=True)
        if adapter._runner:
            await adapter._runner.cleanup()
