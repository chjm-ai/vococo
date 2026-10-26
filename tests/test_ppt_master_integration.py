"""共享上游技能的离线集成验收；未安装时跳过，不联网安装。"""
import os
from pathlib import Path
import subprocess
import sys
import runpy

from pptx import Presentation
import pytest


ROOT = Path(__file__).resolve().parents[1]
ADAPTER = ROOT / "vococo/plugin/skills/editable-pptx"


def test_adapter_uses_upstream_design_and_keeps_host_delivery_boundaries():
    entry = (ADAPTER / "SKILL.md").read_text()
    reference = (ADAPTER / "references/ppt-master.md").read_text()
    assert "默认交给共享安装的 **PPT Master**" in entry
    assert "不是默认流程" in entry
    assert "确认关口必须等待用户" in entry
    assert "VOCOCO_PPT_MASTER_DIR" in reference
    assert "不静默" in reference or "不切换自建生成器" in reference
    assert "--native-charts-and-tables" in reference
    assert "send_file" in reference and "SVG 截图" in reference


def test_pin_accepts_reviewed_tree_and_rejects_missing_or_modified_sources(tmp_path, monkeypatch):
    checker = runpy.run_path(str(ADAPTER / "scripts/ppt_master_check.py"))
    verify = checker["verify"]
    with pytest.raises(ValueError, match="未安装"):
        verify(tmp_path)
    (tmp_path / "SKILL.md").write_text("version: reviewed")
    scripts = tmp_path / "scripts"
    scripts.mkdir()
    (scripts / "attribution_guard.py").write_text("raise SystemExit(0)")
    monkeypatch.setitem(verify.__globals__, "EXPECTED_DIGEST", checker["source_digest"](tmp_path))
    assert verify(tmp_path) == tmp_path / "SKILL.md"
    (tmp_path / "SKILL.md").write_text("version: unreviewed")
    with pytest.raises(ValueError, match="不一致"):
        verify(tmp_path)


def test_pin_rejects_internal_symlink_before_executing_guard(tmp_path):
    checker = runpy.run_path(str(ADAPTER / "scripts/ppt_master_check.py"))
    (tmp_path / "SKILL.md").write_text("test")
    (tmp_path / "alias").symlink_to(tmp_path / "SKILL.md")
    with pytest.raises(ValueError, match="软链接"):
        checker["verify"](tmp_path)


def _run(script, *arguments):
    result = subprocess.run(
        [sys.executable, "-B", str(script), *map(str, arguments)],
        capture_output=True, text=True, timeout=90,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_installed_upstream_exports_native_editable_text_without_network(tmp_path):
    skill = Path(os.environ.get(
        "VOCOCO_PPT_MASTER_DIR", "~/Desktop/Repos/AI_Skills/ppt-master"
    )).expanduser()
    if not (skill / "SKILL.md").is_file():
        pytest.skip("未安装共享 PPT Master；测试不联网安装")
    checker = runpy.run_path(str(ADAPTER / "scripts/ppt_master_check.py"))
    checker["verify"](skill.resolve())
    scripts = skill / "scripts"
    source = tmp_path / "svg_output"
    source.mkdir()
    svg = source / "01_native.svg"
    content = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"
lang="zh-CN" font-family="Microsoft YaHei, Arial" font-size="32" data-pptx-page-role="content">
<rect id="bg" data-pptx-role="background" x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>
<g id="title" data-pptx-bounds="72 72 1136 140" fill="#111111">
<text x="72" y="140">原生文字 Native text</text>
<rect x="72" y="180" width="100" height="4" fill="#2563EB"/>
</g></svg>'''
    svg.write_text(content)
    _run(scripts / "svg_quality_checker.py", tmp_path,
         "--quick-generate", "--canonical-authoring", "--stage", "final", "--json")
    output = tmp_path / "native.pptx"
    _run(scripts / "svg_to_pptx.py", tmp_path,
         "--quick-generate", "--no-notes", "-o", output)
    original = output.read_bytes()
    deck = Presentation(output)
    assert len(deck.slides) == 1
    group = next(shape for shape in deck.slides[0].shapes if hasattr(shape, "shapes"))
    text = next(shape for shape in group.shapes if shape.has_text_frame)
    assert text.text == "原生文字 Native text"
    text.text = "修改后的原生文字"
    edited = tmp_path / "edited.pptx"
    deck.save(edited)
    reopened = Presentation(edited)
    assert "修改后的原生文字" in reopened.slides[0].shapes[0].shapes[0].text
    assert output.read_bytes() == original
    assert svg.read_text() == content
