"""原生 PPT 对象工作台：真实文件创建、局部编辑、稳定ID与版本保护。"""
from io import BytesIO
from pathlib import Path
import importlib.util
import json
import subprocess
import sys
import zipfile

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "vococo/plugin/skills/editable-pptx/scripts/pptx_workspace.py"
spec = importlib.util.spec_from_file_location("pptx_workspace", SCRIPT)
workspace = importlib.util.module_from_spec(spec)
spec.loader.exec_module(workspace)


@pytest.fixture
def deck(tmp_path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(10), Inches(1))
    title.name = "独立标题"
    run = title.text_frame.paragraphs[0].add_run()
    run.text, run.font.size, run.font.bold = "任意原生构图", Pt(32), True
    run.font.color.rgb = RGBColor.from_string("1677FF")
    table = slide.shapes.add_table(2, 2, Inches(.7), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "原始单元格"
    data = CategoryChartData()
    data.categories = ["甲", "乙"]
    data.add_series("测试数据", [20, 40])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2), Inches(5), Inches(3), data)
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(1), Inches(5), Inches(1), Inches(1))
    image = BytesIO()
    Image.new("RGB", (20, 10), "#1677FF").save(image, format="PNG")
    slide.shapes.add_picture(image, Inches(10), Inches(5), width=Inches(1))
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2), Inches(5), Inches(5), Inches(6))
    presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation


def test_inspect_recurses_groups_and_reports_native_objects(deck, tmp_path):
    source = tmp_path / "native.pptx"
    workspace.save_new(deck, source)
    report = workspace.describe(workspace.open_deck(source))
    assert len(report["slides"]) == 2 and report["visual_check"] == "not_run"
    shapes = report["slides"][0]["shapes"]
    assert shapes[0]["paragraphs"][0]["runs"][0]["font"]["size_pt"] == 32
    assert shapes[0]["paragraphs"][0]["runs"][0]["font"]["color"] == "1677FF"
    assert shapes[1]["table"][0][0] == "原始单元格"
    assert shapes[2]["chart"]["series"][0]["values"] == [20, 40]
    assert shapes[3]["children"][0]["parent_ids"] == [shapes[3]["shape_id"]]
    assert shapes[4]["image"]["pixels"] == (20, 10)


def test_edit_existing_objects_keeps_ids_formatting_and_unmodified_picture(deck, tmp_path):
    source, output = tmp_path / "v1.pptx", tmp_path / "v2.pptx"
    workspace.save_new(deck, source)
    original = source.read_bytes()
    picture = deck.slides[0].shapes[4].image.blob
    edited = workspace.open_deck(source)
    slide = workspace.slide_by_id(edited, deck.slides[0].slide_id)
    title = workspace.shape_by_id(slide, deck.slides[0].shapes[0].shape_id)
    title.text_frame.paragraphs[0].runs[0].text = "修改后保留格式"
    title.left = Inches(1)
    slide.shapes[1].table.cell(0, 0).text = "改后数据"
    data = CategoryChartData()
    data.categories = ["甲", "乙"]
    data.add_series("更新系列", [70, 90])
    slide.shapes[2].chart.replace_data(data)
    workspace.move_slide(edited, slide.slide_id, 1)
    receipt = workspace.save_new(edited, output)
    assert receipt["slides"] == 2 and source.read_bytes() == original
    reopened = workspace.open_deck(output)
    assert reopened.slides[1].slide_id == slide.slide_id
    title = workspace.shape_by_id(reopened.slides[1], title.shape_id)
    assert title.text == "修改后保留格式" and title.left == Inches(1)
    assert title.text_frame.paragraphs[0].runs[0].font.bold is True
    assert reopened.slides[1].shapes[1].table.cell(0, 0).text == "改后数据"
    assert list(reopened.slides[1].shapes[2].chart.series[0].values) == [70, 90]
    assert reopened.slides[1].shapes[4].image.blob == picture
    with zipfile.ZipFile(output) as archive:
        assert any(name.endswith(".xlsx") for name in archive.namelist())


def test_add_remove_slides_objects_and_reorder_layers(deck, tmp_path):
    slide = deck.slides[0]
    first = slide.shapes[0].shape_id
    workspace.move_shape(slide, first, len(slide.shapes) - 1)
    assert slide.shapes[-1].shape_id == first
    workspace.move_shape(slide, first, 0)
    assert slide.shapes[0].shape_id == first
    group = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
    child_id = group.shapes[0].shape_id
    assert workspace.shape_by_id(slide, child_id).shape_id == child_id
    workspace.remove_shape(slide, child_id)
    assert len(group.shapes) == 0
    workspace.remove_slide(deck, deck.slides[1].slide_id)
    added = deck.slides.add_slide(deck.slide_layouts[6])
    added.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "新页面"
    output = tmp_path / "operations.pptx"
    workspace.save_new(deck, output)
    assert len(workspace.open_deck(output).slides) == 2


def test_save_refuses_original_existing_file_and_symlink(deck, tmp_path):
    source = tmp_path / "source.pptx"
    workspace.save_new(deck, source)
    original = source.read_bytes()
    with pytest.raises(FileExistsError):
        workspace.save_new(deck, source)
    link = tmp_path / "linked.pptx"
    link.symlink_to(source)
    with pytest.raises(FileExistsError):
        workspace.save_new(deck, link)
    with pytest.raises(ValueError, match="软链接"):
        workspace.open_deck(link)
    assert source.read_bytes() == original


def test_invalid_ids_indices_and_last_slide_fail_explicitly(deck):
    with pytest.raises(ValueError, match="slide_id"):
        workspace.slide_by_id(deck, 999)
    with pytest.raises(ValueError, match="shape_id"):
        workspace.shape_by_id(deck.slides[0], 999)
    with pytest.raises(ValueError, match="页面序号"):
        workspace.move_slide(deck, deck.slides[0].slide_id, 2)
    with pytest.raises(ValueError, match="图层序号"):
        workspace.move_shape(deck.slides[0], deck.slides[0].shapes[0].shape_id, -1)
    workspace.remove_slide(deck, deck.slides[1].slide_id)
    with pytest.raises(ValueError, match="最后一页"):
        workspace.remove_slide(deck, deck.slides[0].slide_id)


def test_page_links_prevent_deleting_navigation_target(deck):
    source, target = deck.slides[0], deck.slides[1]
    source.shapes[0].click_action.target_slide = target
    with pytest.raises(ValueError, match="页内跳转"):
        workspace.remove_slide(deck, target.slide_id)
    assert len(deck.slides) == 2


def test_cli_can_select_stable_slide_id_and_reports_invalid_input(deck, tmp_path):
    source = tmp_path / "native.pptx"
    workspace.save_new(deck, source)
    selected = deck.slides[1].slide_id
    result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--slide-id", str(selected)],
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    assert [slide["slide_id"] for slide in json.loads(result.stdout)["slides"]] == [selected]
    source.write_bytes(b"not a presentation")
    result = subprocess.run([sys.executable, str(SCRIPT), str(source)], capture_output=True, text=True, timeout=30)
    assert result.returncode == 1 and "error" in json.loads(result.stdout)


def test_delete_middle_slide_then_add_preserves_all_slide_contents(tmp_path):
    deck = Presentation()
    for label in ("A", "B", "C"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = label
    workspace.remove_slide(deck, deck.slides[1].slide_id)
    added = deck.slides.add_slide(deck.slide_layouts[6])
    added.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "D"
    output = tmp_path / "delete-middle-add.pptx"
    workspace.save_new(deck, output)
    assert [slide.shapes[0].text for slide in workspace.open_deck(output).slides] == ["A", "C", "D"]
    with zipfile.ZipFile(output) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))


def test_deleting_last_jump_object_releases_slide_relation(deck, tmp_path):
    source, target = deck.slides[0], deck.slides[1]
    first, second = source.shapes[0], source.shapes[1]
    first.click_action.target_slide = second.click_action.target_slide = target
    workspace.remove_shape(source, first.shape_id)
    with pytest.raises(ValueError, match="页内跳转"):
        workspace.remove_slide(deck, target.slide_id)
    workspace.remove_shape(source, second.shape_id)
    workspace.remove_slide(deck, target.slide_id)
    output = tmp_path / "delete-linked.pptx"
    workspace.save_new(deck, output)
    assert len(workspace.open_deck(output).slides) == 1


def test_deleting_picture_retains_media_relation_still_used_by_other_object(deck, tmp_path):
    slide = deck.slides[0]
    first = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    blob = first.image.blob
    second = slide.shapes.add_picture(BytesIO(blob), 0, 0, width=Inches(1))
    workspace.remove_shape(slide, first.shape_id)
    output = tmp_path / "shared-image.pptx"
    workspace.save_new(deck, output)
    reopened = workspace.open_deck(output)
    assert workspace.shape_by_id(reopened.slides[0], second.shape_id).image.blob == blob


def test_open_refuses_oversize_package_and_non_pptx(tmp_path):
    source = tmp_path / "large.pptx"
    with source.open("wb") as output:
        output.truncate(64 * 1024 * 1024 + 1)
    with pytest.raises(ValueError, match="64MB"):
        workspace.open_deck(source)
    wrong = tmp_path / "wrong.txt"
    wrong.write_text("not a deck")
    with pytest.raises(ValueError, match="普通 .pptx"):
        workspace.open_deck(wrong)
