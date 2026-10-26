"""AI 的原生 PPT 工作台：检查对象、稳定定位、版本另存；不规定页面版式。"""
from io import BytesIO
from pathlib import Path
import argparse
import hashlib
import json
import zipfile

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


MAX_BYTES = 64 * 1024 * 1024


def open_deck(source):
    """读取受控工作目录中的稿件；文件权限仍由宿主的工具/沙箱控制。"""
    source = Path(source)
    if source.is_symlink() or not source.is_file() or source.suffix.lower() != ".pptx":
        raise ValueError("输入须为普通 .pptx 文件，不能是软链接。")
    if source.stat().st_size > MAX_BYTES:
        raise ValueError("PPTX 超过 64MB。")
    data = source.read_bytes()
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
            raise ValueError("PPTX 包内容超过检查上限。")
        if any(item.flag_bits & 1 for item in entries):
            raise ValueError("不支持加密 PPTX。")
    return Presentation(BytesIO(data))


def save_new(deck, destination):
    """在内存完成序列化后独占新建，绝不覆盖输入或旧版本。"""
    destination = Path(destination)
    if destination.suffix.lower() != ".pptx":
        raise ValueError("输出须为 .pptx。")
    buffer = BytesIO()
    deck.save(buffer)
    data = buffer.getvalue()
    if len(data) > MAX_BYTES:
        raise ValueError("PPTX 超过 64MB。")
    with destination.open("xb") as output:
        output.write(data)
    return {"output": str(destination), "slides": len(deck.slides), "bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(), "visual_check": "not_run"}


def slide_by_id(deck, slide_id):
    for slide in deck.slides:
        if slide.slide_id == slide_id:
            return slide
    raise ValueError(f"没有 slide_id={slide_id} 的页面。")


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def shape_by_id(slide, shape_id):
    matches = [shape for shape in walk_shapes(slide.shapes) if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"shape_id={shape_id} 命中 {len(matches)} 个对象，不能安全定位。")
    return matches[0]


def remove_shape(slide, shape_id):
    shape = shape_by_id(slide, shape_id)
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    relationships = {value for item in shape._element.iter() for key, value in item.attrib.items()
                     if key.startswith(namespace)}
    shape._element.getparent().remove(shape._element)
    remaining = {value for item in slide._element.iter() for key, value in item.attrib.items()
                 if key.startswith(namespace)}
    for relationship in relationships - remaining:
        if relationship in slide.part.rels:
            slide.part.drop_rel(relationship)


def move_shape(slide, shape_id, index):
    """index 是同组内从底到顶的零基图层序号。"""
    shape = shape_by_id(slide, shape_id)
    parent = shape._element.getparent()
    siblings = [item for item in parent if item.tag.rsplit("}", 1)[-1] in
                {"sp", "pic", "graphicFrame", "cxnSp", "grpSp", "contentPart"}]
    if not 0 <= index < len(siblings):
        raise ValueError("图层序号超出同组对象范围。")
    anchor = siblings[index]
    if anchor is shape._element:
        return
    parent.remove(shape._element)
    if index < siblings.index(shape._element):
        anchor.addprevious(shape._element)
    else:
        anchor.addnext(shape._element)


def move_slide(deck, slide_id, index):
    if not 0 <= index < len(deck.slides):
        raise ValueError("页面序号超出范围。")
    slide_by_id(deck, slide_id)
    entries = deck.slides._sldIdLst
    entry = next(item for item in entries if int(item.id) == slide_id)
    entries.remove(entry)
    entries.insert(index, entry)


def remove_slide(deck, slide_id):
    slide_by_id(deck, slide_id)
    if len(deck.slides) <= 1:
        raise ValueError("不能删除最后一页。")
    # 复杂导航引用不能靠删除页关系无损处理，必须明确失败而非留下坏引用。
    if any(item.tag.rsplit("}", 1)[-1] in {"custShowLst", "sectionLst"} for item in deck._element.iter()):
        raise ValueError("稿件含自定义放映或节，请先显式处理导航引用再删除页面。")
    target = slide_by_id(deck, slide_id).part
    for slide in deck.slides:
        if slide.part is not target and any(rel.target_part is target for rel in slide.part.rels.values() if not rel.is_external):
            raise ValueError("其他页面引用了待删除页，请先处理页内跳转。")
    entry = next(item for item in deck.slides._sldIdLst if int(item.id) == slide_id)
    deck.part.drop_rel(entry.rId)
    deck.slides._sldIdLst.remove(entry)
    # 库的下一页文件名按页数+1分配；删除中间页后须消除编号空洞，避免覆盖仍存活的页。
    deck.part.rename_slide_parts(item.rId for item in deck.slides._sldIdLst)


def _font(font):
    result = {"name": font.name, "size_pt": font.size.pt if font.size else None,
              "bold": font.bold, "italic": font.italic}
    if font.color.type == MSO_COLOR_TYPE.RGB:
        result["color"] = str(font.color.rgb)
    elif font.color.type is not None:
        result["color_type"] = str(font.color.type)
    return result


def _paragraph(paragraph):
    return {"text": paragraph.text, "level": paragraph.level,
            "alignment": str(paragraph.alignment), "font": _font(paragraph.font),
            "runs": [{"text": run.text, "font": _font(run.font)} for run in paragraph.runs]}


def _chart(chart):
    result = {"type": str(chart.chart_type)}
    try:
        result["categories"] = [category.label for category in chart.plots[0].categories]
        result["series"] = [{"name": series.name, "values": list(series.values)} for series in chart.series]
    except (AttributeError, ValueError, NotImplementedError) as exc:
        result["data_warning"] = f"此类图表需要直接检查原生 chart/XML：{exc}"
    return result


def _shape(shape, layer, parents):
    result = {"shape_id": shape.shape_id, "name": shape.name, "type": str(shape.shape_type),
              "layer": layer, "parent_ids": parents, "coordinate_space": "parent" if parents else "slide",
              "box_inches": {key: round(getattr(shape, key) / 914400, 5) for key in ("left", "top", "width", "height")},
              "rotation": shape.rotation}
    if shape.has_text_frame:
        result["text"] = shape.text
        result["paragraphs"] = [_paragraph(item) for item in shape.text_frame.paragraphs]
    if shape.has_table:
        result["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
    if shape.has_chart:
        result["chart"] = _chart(shape.chart)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        result["image"] = {"extension": shape.image.ext, "pixels": shape.image.size,
                           "sha256": hashlib.sha256(shape.image.blob).hexdigest()}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        result["children"] = [_shape(child, i, parents + [shape.shape_id]) for i, child in enumerate(shape.shapes)]
    return result


def describe(deck):
    slides = []
    for index, slide in enumerate(deck.slides):
        slides.append({"index": index, "slide_id": slide.slide_id,
                       "shapes": [_shape(shape, layer, []) for layer, shape in enumerate(slide.shapes)],
                       "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""})
    return {"width_inches": deck.slide_width / 914400, "height_inches": deck.slide_height / 914400,
            "layouts": [{"index": i, "name": layout.name} for i, layout in enumerate(deck.slide_layouts)],
            "slides": slides, "visual_check": "not_run",
            "note": "坐标为英寸，字号为pt；index/layer从0开始；继承字体可能为null，不代表没有样式。"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--slide-id", type=int, help="只输出稳定ID对应的页面")
    args = parser.parse_args()
    try:
        deck = open_deck(args.pptx)
        report = describe(deck)
        if args.slide_id is not None:
            slide_by_id(deck, args.slide_id)
            report["slides"] = [slide for slide in report["slides"] if slide["slide_id"] == args.slide_id]
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
